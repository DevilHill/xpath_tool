"""基于 openpyxl 的离线 XLSX 子集，不依赖或控制 Excel/WPS 进程。

需要可选依赖 openpyxl。单元格和工作表序号均从 1 开始；range 使用 A1 表达式。
"""
from pathlib import Path


def _backend():
    try:
        import openpyxl
    except ImportError as error:
        raise ImportError("Excel 功能需要 openpyxl，请运行 pip install 'wbot[excel]' 或 pip install openpyxl") from error
    return openpyxl


def _check_kind(kind):
    if kind not in ("openpyxl", "xlsx"):
        raise NotImplementedError("wbot.excel 仅支持 kind='openpyxl' 或 'xlsx'，不控制 Office/WPS")


def _positive(value, name):
    if not isinstance(value, int) or isinstance(value, bool) or value < 1:
        raise ValueError(f"{name} 必须为从 1 开始的正整数")
    return value


class WorkRange:
    def __init__(self, sheet, address):
        self.sheet = sheet
        self.address = address
        bounds = _backend().utils.range_boundaries(address)
        if any(value is None for value in bounds):
            raise ValueError("range 必须为有限的 A1 或 A1:C8 区域")
        self.min_col, self.min_row, self.max_col, self.max_row = bounds
        if not (1 <= self.min_col <= self.max_col <= 16384 and 1 <= self.min_row <= self.max_row <= 1048576):
            raise ValueError("range 超出 Excel 单元格边界")

    @property
    def value(self):
        self.sheet._check()
        rows = list(self.sheet._sheet.iter_rows(min_row=self.min_row, max_row=self.max_row,
                    min_col=self.min_col, max_col=self.max_col, values_only=True))
        return rows[0][0] if self.min_row == self.max_row and self.min_col == self.max_col else [list(row) for row in rows]

    @value.setter
    def value(self, values):
        self.sheet._check()
        if self.min_row == self.max_row and self.min_col == self.max_col:
            self.sheet._sheet.cell(self.min_row, self.min_col).value = values
            return
        rows = [list(row) for row in values]
        width, height = self.max_col - self.min_col + 1, self.max_row - self.min_row + 1
        if len(rows) != height or any(len(row) != width for row in rows):
            raise ValueError("二维数据尺寸必须与目标区域一致")
        for row_number, row in enumerate(rows, self.min_row):
            for col_number, value in enumerate(row, self.min_col):
                self.sheet._sheet.cell(row_number, col_number).value = value

    def get_value(self):
        return self.value

    def set_value(self, value):
        self.value = value
        return self

    def clear(self):
        self.sheet._check()
        for row in self.sheet._sheet.iter_rows(min_row=self.min_row, max_row=self.max_row, min_col=self.min_col, max_col=self.max_col):
            for cell in row:
                cell.value = None


class WorkSheet:
    def __init__(self, workbook, sheet):
        self.workbook, self._sheet = workbook, sheet

    def _check(self):
        self.workbook._check()

    @property
    def name(self):
        self._check()
        return self._sheet.title

    @name.setter
    def name(self, value):
        self._check()
        if value != self._sheet.title and value in self.workbook.sheet_names:
            raise ValueError("工作表名称已存在")
        self._sheet.title = value

    def cell(self, column_name, row_num):
        self._check()
        _positive(row_num, "row_num")
        column = _backend().utils.get_column_letter(_positive(column_name, "column_name")) if isinstance(column_name, int) else column_name.upper()
        return WorkRange(self, f"{column}{row_num}")

    def range(self, destination):
        self._check()
        return WorkRange(self, destination)

    def append_row(self, values, begin_column_name="A"):
        self._check()
        values = list(values)
        col = _backend().utils.column_index_from_string(begin_column_name)
        # 新建空工作表的 max_row 为 1，先检查该行是否为空。
        row = 1 if self._sheet.max_row == 1 and all(cell.value is None for cell in self._sheet[1]) else self._sheet.max_row + 1
        for offset, value in enumerate(values):
            self._sheet.cell(row, col + offset).value = value
        return row

    def append_column(self, values, begin_row_index=1):
        self._check()
        _positive(begin_row_index, "begin_row_index")
        col = 1 if self._sheet.max_column == 1 and all(row[0].value is None for row in self._sheet.iter_rows()) else self._sheet.max_column + 1
        for offset, value in enumerate(values):
            self._sheet.cell(begin_row_index + offset, col).value = value
        return col

    def read(self, address=None):
        self._check()
        return self.range(address).value if address else [list(row) for row in self._sheet.iter_rows(values_only=True)]

    def write(self, address, values):
        self.range(address).value = values


class WorkBook:
    def __init__(self, workbook, path=None):
        self._workbook, self.path, self._closed = workbook, Path(path) if path else None, False

    def _check(self):
        if self._closed:
            raise ValueError("工作簿已关闭")

    @property
    def sheet_names(self):
        self._check()
        return list(self._workbook.sheetnames)

    def get_sheet_by_index(self, index):
        self._check()
        _positive(index, "index")
        return WorkSheet(self, self._workbook.worksheets[index - 1])

    def get_sheet_by_name(self, name):
        self._check()
        return WorkSheet(self, self._workbook[name])

    def get_active_sheet(self):
        self._check()
        return WorkSheet(self, self._workbook.active)

    def create_sheet(self, name, create_way="last"):
        self._check()
        if name in self.sheet_names:
            raise ValueError("工作表名称已存在")
        if create_way not in ("first", "last"):
            raise ValueError("create_way 必须为 first 或 last")
        return WorkSheet(self, self._workbook.create_sheet(name, 0 if create_way == "first" else None))

    def delete_sheet(self, name):
        self._check()
        if len(self.sheet_names) < 2:
            raise ValueError("至少保留一个工作表")
        self._workbook.remove(self._workbook[name])

    def save(self):
        self._check()
        if self.path is None:
            raise ValueError("新建工作簿请先调用 save_as(path)")
        return self.save_as(self.path, overwrite=True)

    def save_as(self, filename, *, overwrite=False):
        self._check()
        target = Path(filename)
        if target.suffix.lower() not in (".xlsx", ".xlsm"):
            raise ValueError("仅支持 .xlsx 或 .xlsm 文件")
        with target.open("wb" if overwrite else "xb") as stream:
            self._workbook.save(stream)
        self.path = target
        return str(target)

    def close(self):
        if not self._closed:
            self._workbook.close()
            self._closed = True

    def __enter__(self):
        self._check()
        return self

    def __exit__(self, *exc):
        self.close()


def create(*, kind="openpyxl", visible=False, original_file=""):
    _check_kind(kind)
    if visible:
        raise NotImplementedError("离线 Excel 后端不支持 visible=True")
    if original_file:
        workbook = open(original_file, kind=kind)
        workbook.path = None
        return workbook
    return WorkBook(_backend().Workbook())


def open(file_name, *, kind="openpyxl", visible=False, password="", write_password="", ignore_formula=False, update_links=False):
    _check_kind(kind)
    if visible or password or write_password or update_links:
        raise NotImplementedError("离线后端不支持可视窗口、加密密码或刷新外部链接")
    target = Path(file_name)
    if target.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("仅支持 .xlsx 或 .xlsm 文件")
    workbook = _backend().load_workbook(target, data_only=ignore_formula, keep_vba=target.suffix.lower() == ".xlsm", keep_links=False)
    return WorkBook(workbook, target)
